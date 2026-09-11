"""生成**合成**演示数据：规则表 + 综测表 + 往年汇总（ground truth）。

重要：这里生成的每一行都是编造的，不含任何真实同学信息 —— 这个文件可以安全提交。
真实材料请一律放到 ``data/raw/``（已被 .gitignore 隔离）。

用法:
    uv run python scripts/make_sample_data.py            # 默认写 data/sample/
    uv run python scripts/make_sample_data.py --out demo # 指定输出目录
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent

# ======================================================================
# 合成规则：模拟一份"等级 × 类别 加分矩阵"
# 故意包含合并单元格，用来演示 fill-down 为什么是必需的
# ======================================================================
RULE_ROWS: list[tuple[str | None, str, float, float | None]] = [
    ("学科竞赛", "国家级一等奖", 30, 40),
    (None, "国家级二等奖", 25, 40),
    (None, "省级一等奖", 15, 30),
    (None, "省级二等奖", 12, 30),
    (None, "省级三等奖", 8, 30),
    (None, "校级一等奖", 5, 15),
    ("文体活动", "一等奖", 6, 12),
    (None, "二等奖", 4, 12),
    (None, "三等奖", 2, 12),
    ("荣誉称号", "国家奖学金", 10, 10),
    (None, "省级三好学生", 5, 10),
    ("科研成果", "发明专利", 20, 30),
    (None, "软件著作权", 8, 30),
]

# 学生 1 的三条学科竞赛不累加（同类取最高），刻意用来演示"取最高"
CLAIM_ROWS: list[tuple[str, str, str, str, str, str]] = [
    ("2023001", "张三", "学科竞赛", "省二等奖", "省级二等奖", "否"),
    ("2023001", "张三", "学科竞赛", "国赛二等奖", "国家级二等奖", "否"),
    ("2023001", "张三", "学科竞赛", "校赛一等奖", "校级一等奖", "否"),
    ("2023001", "张三", "文体活动", "校运会一等奖", "一等奖", "否"),
    ("2023002", "李四", "学科竞赛", "省赛三等奖", "省级三等奖", "是"),
    ("2023002", "李四", "荣誉称号", "国家奖学金", "国家奖学金", "否"),
    ("2023003", "王五", "科研成果", "发明专利一项", "发明专利", "否"),
    ("2023003", "王五", "科研成果", "软著一项", "软件著作权", "否"),
    ("2023003", "王五", "学科竞赛", "省一等奖", "省级一等奖", "是"),
    ("2023004", "赵六", "学科竞赛", "参加了数学建模比赛", None, "否"),
    ("2023004", "赵六", "文体活动", "校运会二等奖", "二等奖", "否"),
    ("2023005", "钱七", "宿舍管理", "宿舍卫生优秀", None, "否"),
    ("2023005", "钱七", "文体活动", "校运会三等奖", "三等奖", "否"),
]


def write_rules(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "加分标准表"
    ws.append(["类别", "等级", "分值", "封顶"])
    start: int | None = None  # 当前类别块的起始行；下一个类别出现时合并
    for category, level, score, cap in RULE_ROWS:
        if category is not None:
            if start is not None and ws.max_row > start:
                ws.merge_cells(start_row=start, start_column=1, end_row=ws.max_row, end_column=1)
            start = ws.max_row + 1
            ws.append([category, level, score, cap])
        else:
            ws.append([None, level, score, cap])
    if start is not None and ws.max_row > start:
        ws.merge_cells(start_row=start, start_column=1, end_row=ws.max_row, end_column=1)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 18
    wb.save(path)
    print(f"[rules] {path.name}: {len(RULE_ROWS)} 条（含合并单元格）")


def write_claims(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "综测汇总"
    ws.append(["学号", "姓名", "类别", "申报内容", "等级", "是否团队"])
    for row in CLAIM_ROWS:
        ws.append(list(row))
    # 同一个人的多条申报：学号/姓名列合并（真实综测表几乎都这么写）
    block_start = 2  # 第 1 行是表头
    for i in range(1, len(CLAIM_ROWS) + 1):
        current = CLAIM_ROWS[i - 1][0]
        nxt = CLAIM_ROWS[i][0] if i < len(CLAIM_ROWS) else None
        if nxt != current:
            end = i + 1  # CLAIM_ROWS[i-1] 对应工作表第 i+1 行
            if block_start != end:
                ws.merge_cells(start_row=block_start, start_column=1, end_row=end, end_column=1)
                ws.merge_cells(start_row=block_start, start_column=2, end_row=end, end_column=2)
            block_start = end + 1
    wb.save(path)
    print(f"[claims] {path.name}: {len(CLAIM_ROWS)} 条申报，"
          f"{len({r[0] for r in CLAIM_ROWS})} 名学生（学号列含合并单元格）")


def write_truth(path: Path, truth: dict[str, float], names: dict[str, str]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "汇总"
    ws.append(["学号", "姓名", "总分"])
    for sid, total in sorted(truth.items()):
        ws.append([sid, names.get(sid, ""), total])
    wb.save(path)
    print(f"[truth]  {path.name}: {len(truth)} 人（由引擎生成，供回测演示）")


def main() -> int:
    parser = argparse.ArgumentParser(description="生成合成演示数据（无真实隐私信息）")
    parser.add_argument("--out", default="data/sample", help="输出目录，默认 data/sample")
    args = parser.parse_args()

    out = (ROOT / args.out) if not Path(args.out).is_absolute() else Path(args.out)
    out.mkdir(parents=True, exist_ok=True)

    rules_path = out / "rules_sample.xlsx"
    claims_path = out / "claims_sample.xlsx"
    truth_path = out / "truth_sample.xlsx"

    write_rules(rules_path)
    write_claims(claims_path)

    # 用真实引擎算一遍作为 ground truth（脚本自己先跑通，才有资格生成"标准答案"）
    from scoreproof.calc.engine import compute_all
    from scoreproof.ingest.excel_loader import load_claims, load_rules
    from scoreproof.schema import Ruleset

    rules = load_rules(rules_path, academic_year="2025-2026")
    claims = load_claims(claims_path, academic_year="2025-2026")
    results = compute_all(claims, Ruleset(rules=rules), academic_year="2025-2026")
    truth = {sid: bd.total for sid, bd in results.items()}
    names = {c.student_id: c.student_name or "" for c in claims}
    write_truth(truth_path, truth, names)

    print("\n期望总分（由计算引擎生成）:")
    for sid, total in sorted(truth.items()):
        bd = results[sid]
        flag = ""
        if bd.unmatched_claims:
            flag += f" [{len(bd.unmatched_claims)} 条未命中]"
        if bd.review_claims:
            flag += f" [{len(bd.review_claims)} 条待复核]"
        print(f"  {sid} {names.get(sid, ''):<4} {total:>6g}{flag}")

    print("\n下一步:")
    print(f"  uv run scoreproof import-rules {args.out}/rules_sample.xlsx --year 2025-2026")
    print(f"  uv run scoreproof calc {args.out}/claims_sample.xlsx --year 2025-2026")
    print(f"  uv run scoreproof backtest {args.out}/claims_sample.xlsx "
          f"--truth {args.out}/truth_sample.xlsx --year 2025-2026")
    return 0


if __name__ == "__main__":
    sys.exit(main())
