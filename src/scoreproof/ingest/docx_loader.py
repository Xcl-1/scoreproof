"""Word 接入：正文 + 表格，**按文档原始顺序**输出。

坑点（项目总结第 2 节）：很多通知把补充规定写在 ``document.tables`` 里，
只读 ``document.paragraphs`` 会直接漏掉整条规定，因此必须遍历 body 元素。
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Literal

from ..errors import DataSourceError

BlockKind = Literal["paragraph", "table"]


@dataclass
class DocxBlock:
    """文档中的一个顺序块（段落或表格）。"""

    index: int
    kind: BlockKind
    text: str = ""
    rows: list[list[str]] | None = None
    style: str | None = None

    @property
    def is_table(self) -> bool:
        return self.kind == "table"

    def as_text(self) -> str:
        if self.kind == "paragraph":
            return self.text
        rows = self.rows or []
        return "\n".join(" | ".join(r) for r in rows)


@dataclass
class DocxDocument:
    path: Path
    blocks: list[DocxBlock]

    @property
    def doc_name(self) -> str:
        return self.path.name

    @property
    def text(self) -> str:
        return "\n".join(b.as_text() for b in self.blocks)

    @property
    def tables(self) -> list[list[list[str]]]:
        return [b.rows or [] for b in self.blocks if b.is_table]

    @property
    def paragraphs(self) -> list[str]:
        return [b.text for b in self.blocks if b.kind == "paragraph" and b.text]

    def find(self, *keywords: str) -> list[DocxBlock]:
        return [b for b in self.blocks if any(k in b.as_text() for k in keywords)]

    def source_ref(self, block: DocxBlock, *, clause: str | None = None, table: str | None = None):
        """块 -> SourceRef（Word 没有页码，用块序号定位 + 原文片段）。"""
        from ..schema import SourceRef

        return SourceRef(
            doc=self.doc_name,
            table=table,
            clause=clause,
            row=block.index + 1,
            text=block.as_text()[:200] or None,
        )

    def meta(self) -> dict[str, Any]:
        return {"doc": self.doc_name, "blocks": len(self.blocks),
                "paragraphs": len(self.paragraphs), "tables": len(self.tables)}


def load_docx(path: str | Path) -> DocxDocument:
    """读取 .docx，保持段落与表格的原始先后顺序。"""
    p = Path(path)
    if not p.exists():
        raise DataSourceError(f"Word 文件不存在：{p}", detail={"path": str(p)})
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover
        raise DataSourceError(
            "需要 python-docx 才能读取 .docx",
            detail={"hint": "uv add python-docx"},
        ) from exc

    document = docx.Document(str(p))
    blocks: list[DocxBlock] = []
    for idx, child in enumerate(document.element.body.iterchildren()):
        tag = child.tag.split("}")[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if text:
                blocks.append(
                    DocxBlock(index=idx, kind="paragraph", text=text,
                              style=getattr(para.style, "name", None))
                )
        elif tag == "tbl":
            table = Table(child, document)
            rows = [
                [cell.text.strip().replace("\n", " ") for cell in row.cells]
                for row in table.rows
            ]
            blocks.append(DocxBlock(index=idx, kind="table", rows=rows))
    return DocxDocument(path=p, blocks=blocks)


def table_to_records(rows: list[list[str]]) -> list[dict[str, str]]:
    """把 Word 表格转成 ``[{表头: 值}]``（首行作表头），便于直接映射成规则。"""
    if not rows:
        return []
    header = [h or f"col_{i}" for i, h in enumerate(rows[0])]
    records: list[dict[str, str]] = []
    for row in rows[1:]:
        if not any(str(c).strip() for c in row):
            continue
        records.append({h: (row[i] if i < len(row) else "") for i, h in enumerate(header)})
    return records


__all__ = ["DocxBlock", "DocxDocument", "load_docx", "table_to_records"]
